"""
Quick test script to verify export libraries are working.
Run: python test_export_direct.py
"""
import io
import sys

print(f"Python: {sys.version}")

# Test 1: openpyxl
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test"
    headers = ["wound_type", "title", "question_text", "answer_text", "triage_level"]
    header_fill = PatternFill(fgColor="17805f", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
    ws.append(["burn", "Test Q", "Question 1", "Answer 1", "green"])
    buf = io.BytesIO()
    wb.save(buf)
    print(f"✅ openpyxl OK: {openpyxl.__version__}, Excel size={len(buf.getvalue())} bytes")
except Exception as e:
    print(f"❌ openpyxl ERROR: {e}")
    import traceback; traceback.print_exc()

# Test 2: CSV
try:
    import csv
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["wound_type", "title", "question_text", "answer_text", "triage_level"])
    writer.writerow(["burn", "Test", "Q1", "A1", "green"])
    result = output.getvalue().encode("utf-8-sig")
    print(f"✅ CSV OK: {len(result)} bytes")
except Exception as e:
    print(f"❌ CSV ERROR: {e}")

# Test 3: reportlab
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    styles = getSampleStyleSheet()
    story = [Paragraph("Test PDF", styles["Title"])]
    doc.build(story)
    import reportlab
    print(f"✅ reportlab OK: {reportlab.Version}, PDF size={len(buf.getvalue())} bytes")
except Exception as e:
    print(f"❌ reportlab ERROR: {e}")
    import traceback; traceback.print_exc()

# Test 4: python-docx
try:
    from docx import Document
    doc = Document()
    doc.add_heading("Test DOCX", level=1)
    buf = io.BytesIO()
    doc.save(buf)
    print(f"✅ python-docx OK: DOCX size={len(buf.getvalue())} bytes")
except Exception as e:
    print(f"❌ python-docx ERROR: {e}")

print("\nDone.")
