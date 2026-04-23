from __future__ import annotations

from pathlib import Path

import fitz  # PyMuPDF
from langchain_community.document_loaders import (
    BSHTMLLoader,
    CSVLoader,
    TextLoader,
)

from app.core.config import settings
from app.modules.rag.exceptions import LoaderError, UnsupportedFileType

_SUPPORTED = {"pdf", "md", "txt", "docx", "html", "csv"}
_OCR_LANG = "vie+eng"
_OCR_DPI = 300


class LoaderService:
    def load(self, path: str, file_type: str) -> str:
        ext = file_type.lower().lstrip(".")
        if ext not in _SUPPORTED:
            raise UnsupportedFileType(f"File type '{ext}' không hỗ trợ")

        file_path = Path(path)
        if not file_path.exists():
            raise LoaderError(f"File không tồn tại: {path}")

        try:
            text = self._extract(ext, str(file_path))
            if not text.strip():
                raise LoaderError(f"File rỗng sau khi load: {path}")
            return text
        except (UnsupportedFileType, LoaderError):
            raise
        except Exception as exc:
            raise LoaderError(f"Không load được file {path}: {exc}") from exc

    def _extract(self, ext: str, path: str) -> str:
        if ext == "pdf":
            return self._load_pdf(path)
        if ext == "docx":
            return self._load_docx(path)
        if ext in ("md", "txt"):
            loader = TextLoader(path, encoding="utf-8", autodetect_encoding=True)
        elif ext == "html":
            loader = BSHTMLLoader(path, open_encoding="utf-8")
        elif ext == "csv":
            loader = CSVLoader(path, encoding="utf-8")
        else:
            raise UnsupportedFileType(ext)
        docs = loader.load()
        return "\n\n".join(d.page_content for d in docs if d.page_content)

    def _load_docx(self, path: str) -> str:
        from docx import Document  # python-docx

        doc = Document(path)
        parts: list[str] = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n\n".join(parts)

    def _load_pdf(self, path: str) -> str:
        text_parts: list[str] = []
        with fitz.open(path) as pdf:
            for page in pdf:
                text_parts.append(page.get_text("text"))
        text = "\n\n".join(t for t in text_parts if t.strip())
        if text.strip():
            return text
        return self._ocr_pdf(path)

    def _ocr_pdf(self, path: str) -> str:
        try:
            import pytesseract
            from PIL import Image
        except ImportError as exc:
            raise LoaderError(
                f"OCR không khả dụng — cài pytesseract + Pillow: {exc}"
            ) from exc

        if settings.TESSERACT_CMD:
            pytesseract.pytesseract.tesseract_cmd = settings.TESSERACT_CMD

        zoom = _OCR_DPI / 72
        matrix = fitz.Matrix(zoom, zoom)
        pages: list[str] = []
        with fitz.open(path) as pdf:
            for page in pdf:
                pix = page.get_pixmap(matrix=matrix, alpha=False)
                img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                try:
                    pages.append(pytesseract.image_to_string(img, lang=_OCR_LANG))
                except pytesseract.TesseractNotFoundError as exc:
                    raise LoaderError(
                        "Tesseract binary không tìm thấy — cài Tesseract OCR và set TESSERACT_CMD trong .env"
                    ) from exc
                except Exception:
                    pass

        return "\n\n".join(p for p in pages if p.strip())

