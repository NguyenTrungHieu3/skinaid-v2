"""
Export Service for Questionnaire module.
- Export: DOCX (.docx) / PDF (.pdf via reportlab) / CSV / Excel (re-importable full format)
"""
import io
import csv
import logging

from app.modules.questionnaires.models.questionnaire import Questionnaire

logger = logging.getLogger(__name__)

TRIAGE_COLORS = {
    "green":  (0.18, 0.80, 0.44),
    "yellow": (1.0, 0.80, 0.0),
    "red":    (0.9, 0.20, 0.20),
}
TRIAGE_LABELS = {"green": "Nhẹ", "yellow": "Vừa", "red": "Nặng"}


# ─── DOCX Export ──────────────────────────────────────────────────────────────

def _render_docx_questionnaire(doc, questionnaire: Questionnaire):
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    # Title
    title_para = doc.add_heading(questionnaire.title, level=1)
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_para.runs[0]
    run.font.color.rgb = RGBColor(0x17, 0x80, 0x5f)

    # Meta info
    status_text = "✅ Đang hoạt động" if questionnaire.is_active else "📝 Bản nháp"
    meta = doc.add_paragraph()
    meta.add_run("Loại vết thương: ").bold = True
    meta.add_run(questionnaire.wound_type)
    meta.add_run("   |   ")
    meta.add_run("Trạng thái: ").bold = True
    meta.add_run(status_text)
    if questionnaire.description:
        desc = doc.add_paragraph()
        desc.add_run("Mô tả: ").bold = True
        desc.add_run(questionnaire.description)

    doc.add_paragraph()  # spacer

    questions = sorted(questionnaire.questions or [], key=lambda q: q.order_index)
    for idx, q in enumerate(questions, start=1):
        # Question heading
        q_para = doc.add_paragraph()
        run = q_para.add_run(f"Câu {idx}: {q.question_text}")
        run.bold = True
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x0f, 0x17, 0x2a)

        mc_label = "(Chọn nhiều)" if q.is_multiple_choice else "(Chọn một)"
        q_para.add_run(f"  {mc_label}").italic = True

        # Answers table
        answers = sorted(q.answers or [], key=lambda a: a.order_index)
        if answers:
            table = doc.add_table(rows=1, cols=2)
            table.style = "Table Grid"
            hdr = table.rows[0].cells
            hdr[0].text = "Đáp án"
            hdr[1].text = "Mức độ Triage"

            for hdr_cell in hdr:
                for p in hdr_cell.paragraphs:
                    for r in p.runs:
                        r.bold = True

            for ans in answers:
                row = table.add_row().cells
                row[0].text = ans.answer_text
                triage_label = TRIAGE_LABELS.get(ans.triage_level, ans.triage_level)
                row[1].text = triage_label

        doc.add_paragraph()

    # Footer
    footer = doc.add_paragraph(f"Xuất bởi SkinAid Admin — {questionnaire.questionnaire_id}")
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = RGBColor(0x94, 0xa3, 0xb8)

def export_to_docx(questionnaire: Questionnaire) -> bytes:
    """Generate a Word (.docx) document for the questionnaire."""
    from docx import Document
    doc = Document()
    _render_docx_questionnaire(doc, questionnaire)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def export_bulk_to_docx(questionnaires: list[Questionnaire]) -> bytes:
    """Generate a single Word (.docx) document containing multiple questionnaires."""
    from docx import Document
    doc = Document()
    for i, questionnaire in enumerate(questionnaires):
        if i > 0:
            doc.add_page_break()
        _render_docx_questionnaire(doc, questionnaire)
        
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ─── PDF Export ───────────────────────────────────────────────────────────────

def _setup_pdf_fonts():
    import os
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    module_dir = os.path.dirname(os.path.dirname(__file__))
    fonts_dir = os.path.join(module_dir, "fonts")
    roboto_path = os.path.join(fonts_dir, "Roboto-Regular.ttf")
    roboto_bold_path = os.path.join(fonts_dir, "Roboto-Bold.ttf")

    if os.path.exists(roboto_path) and os.path.exists(roboto_bold_path):
        pdfmetrics.registerFont(TTFont('Roboto', roboto_path))
        pdfmetrics.registerFont(TTFont('Roboto-Bold', roboto_bold_path))
        pdfmetrics.registerFontFamily('Roboto', normal='Roboto', bold='Roboto-Bold')
        return 'Roboto', 'Roboto-Bold'
    else:
        return 'Helvetica', 'Helvetica-Bold'

def _create_pdf_styles(font_name, font_bold):
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.enums import TA_CENTER
    PRIMARY = colors.HexColor("#17805f")
    styles = getSampleStyleSheet()
    return {
        "title": ParagraphStyle("Title", parent=styles["Title"], fontName=font_bold, textColor=PRIMARY, fontSize=22, spaceAfter=6, alignment=TA_CENTER),
        "meta": ParagraphStyle("Meta", parent=styles["Normal"], fontName=font_name, fontSize=10, textColor=colors.HexColor("#475569"), spaceAfter=4),
        "question": ParagraphStyle("Question", parent=styles["Normal"], fontSize=12, fontName=font_bold, textColor=colors.HexColor("#0f172a"), spaceBefore=12, spaceAfter=6),
        "answer": ParagraphStyle("Answer", parent=styles["Normal"], fontName=font_name, fontSize=10, textColor=colors.HexColor("#334155")),
        "footer": ParagraphStyle("Footer", parent=styles["Normal"], fontName=font_name, fontSize=8, textColor=colors.HexColor("#94a3b8"), alignment=TA_CENTER)
    }

def _render_pdf_questionnaire(story, questionnaire: Questionnaire, styles_dict, font_name, font_bold):
    from reportlab.lib.units import cm
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Spacer, Table, TableStyle, HRFlowable

    PRIMARY = colors.HexColor("#17805f")
    LIGHT_GREEN = colors.HexColor("#dcfce7")
    LIGHT_YELLOW = colors.HexColor("#fef9c3")
    LIGHT_RED = colors.HexColor("#fee2e2")

    TRIAGE_BG = {"green": LIGHT_GREEN, "yellow": LIGHT_YELLOW, "red": LIGHT_RED}
    TRIAGE_FG = {
        "green":  colors.HexColor("#166534"),
        "yellow": colors.HexColor("#854d0e"),
        "red":    colors.HexColor("#991b1b"),
    }

    story.append(Paragraph(questionnaire.title, styles_dict["title"]))
    story.append(HRFlowable(width="100%", thickness=2, color=PRIMARY))
    story.append(Spacer(1, 0.3*cm))

    status_text = "✅ Đang hoạt động" if questionnaire.is_active else "📝 Bản nháp"
    story.append(Paragraph(
        f"<b>Loại vết thương:</b> {questionnaire.wound_type}   |   "
        f"<b>Trạng thái:</b> {status_text}",
        styles_dict["meta"]
    ))
    if questionnaire.description:
        story.append(Paragraph(f"<b>Mô tả:</b> {questionnaire.description}", styles_dict["meta"]))

    story.append(Spacer(1, 0.5*cm))

    questions = sorted(questionnaire.questions or [], key=lambda q: q.order_index)
    for idx, q in enumerate(questions, start=1):
        mc_label = "(Chọn nhiều)" if q.is_multiple_choice else "(Chọn một)"
        story.append(Paragraph(f"Câu {idx}: {q.question_text} <i>{mc_label}</i>", styles_dict["question"]))

        answers = sorted(q.answers or [], key=lambda a: a.order_index)
        if answers:
            table_data = [["Đáp án", "Mức độ Triage"]]
            for ans in answers:
                table_data.append([
                    Paragraph(ans.answer_text, styles_dict["answer"]),
                    Paragraph(TRIAGE_LABELS.get(ans.triage_level, ans.triage_level), styles_dict["answer"])
                ])

            col_widths = [12*cm, 4*cm]
            t = Table(table_data, colWidths=col_widths)
            triage_style = [
                ('BACKGROUND', (0, 0), (-1, 0), PRIMARY),
                ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
                ('FONTNAME', (0, 0), (-1, 0), font_bold),
                ('FONTSIZE', (0, 0), (-1, 0), 10),
                ('ROWBACKGROUND', (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
                ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
                ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
                ('FONTNAME', (0, 1), (-1, -1), font_name),
                ('LEFTPADDING', (0, 0), (-1, -1), 8),
                ('RIGHTPADDING', (0, 0), (-1, -1), 8),
                ('TOPPADDING', (0, 0), (-1, -1), 6),
                ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
            ]
            for row_i, ans in enumerate(answers, start=1):
                bg = TRIAGE_BG.get(ans.triage_level, colors.white)
                fg = TRIAGE_FG.get(ans.triage_level, colors.black)
                triage_style.append(('BACKGROUND', (1, row_i), (1, row_i), bg))
                triage_style.append(('TEXTCOLOR', (1, row_i), (1, row_i), fg))

            t.setStyle(TableStyle(triage_style))
            story.append(t)

        story.append(Spacer(1, 0.3*cm))

    # Footer
    story.append(Spacer(1, 1*cm))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#e2e8f0")))
    story.append(Paragraph(f"Xuất bởi SkinAid Admin | ID: {questionnaire.questionnaire_id}", styles_dict["footer"]))


def export_to_pdf(questionnaire: Questionnaire) -> bytes:
    """Generate a PDF document for the questionnaire using ReportLab."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate
    
    font_name, font_bold = _setup_pdf_fonts()
    styles_dict = _create_pdf_styles(font_name, font_bold)
    
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2*cm, rightMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)
    story = []
    _render_pdf_questionnaire(story, questionnaire, styles_dict, font_name, font_bold)
    doc.build(story)
    return buf.getvalue()

def export_bulk_to_pdf(questionnaires: list[Questionnaire]) -> bytes:
    """Generate a single PDF document containing multiple questionnaires."""
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, PageBreak
    
    font_name, font_bold = _setup_pdf_fonts()
    styles_dict = _create_pdf_styles(font_name, font_bold)
    
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2*cm, rightMargin=2*cm, topMargin=2*cm, bottomMargin=2*cm)
    story = []
    for i, questionnaire in enumerate(questionnaires):
        if i > 0:
            story.append(PageBreak())
        _render_pdf_questionnaire(story, questionnaire, styles_dict, font_name, font_bold)
        
    doc.build(story)
    return buf.getvalue()


# ─── Bulk CSV / Excel Export ──────────────────────────────────────────────────

def export_bulk_to_csv(questionnaires: list[Questionnaire]) -> bytes:
    """
    Export multiple questionnaires as a single CSV in the full-import format.
    The resulting file can be re-imported via /import/bulk or /import/bulk-files.
    """
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow([
        "wound_type", "title", "description", "is_active",
        "question_order", "question_text", "is_multiple_choice",
        "answer_text", "triage_level"
    ])

    for questionnaire in questionnaires:
        questions = sorted(questionnaire.questions or [], key=lambda q: q.order_index)
        if not questions:
            logger.warning("[EXPORT CSV] Questionnaire '%s' has no questions loaded.", questionnaire.title)
            
        for q in questions:
            answers = sorted(q.answers or [], key=lambda a: a.order_index)
            if answers:
                for ans in answers:
                    writer.writerow([
                        questionnaire.wound_type, questionnaire.title, questionnaire.description or "",
                        str(questionnaire.is_active).lower(), q.order_index, q.question_text,
                        str(q.is_multiple_choice).lower(), ans.answer_text, ans.triage_level,
                    ])
            else:
                writer.writerow([
                    questionnaire.wound_type, questionnaire.title, questionnaire.description or "",
                    str(questionnaire.is_active).lower(), q.order_index, q.question_text,
                    str(q.is_multiple_choice).lower(), "", "green",
                ])

    return output.getvalue().encode("utf-8-sig")


def export_bulk_to_excel(questionnaires: list[Questionnaire]) -> bytes:
    """
    Export multiple questionnaires as a single Excel (.xlsx) in the full-import format.
    The resulting file can be re-imported via /import/bulk or /import/bulk-files.
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bộ câu hỏi"

    headers = [
        "wound_type", "title", "description", "is_active",
        "question_order", "question_text", "is_multiple_choice",
        "answer_text", "triage_level"
    ]
    header_fill = PatternFill(fgColor="17805f", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")

    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    triage_fills = {
        "green":  PatternFill(fgColor="dcfce7", fill_type="solid"),
        "yellow": PatternFill(fgColor="fef9c3", fill_type="solid"),
        "red":    PatternFill(fgColor="fee2e2", fill_type="solid"),
    }

    for questionnaire in questionnaires:
        questions = sorted(questionnaire.questions or [], key=lambda q: q.order_index)
        if not questions:
            logger.warning("[EXPORT EXCEL] Questionnaire '%s' has no questions loaded.", questionnaire.title)
            
        for q in questions:
            answers = sorted(q.answers or [], key=lambda a: a.order_index)
            if answers:
                for ans in answers:
                    row = [
                        questionnaire.wound_type, questionnaire.title, questionnaire.description or "",
                        str(questionnaire.is_active).lower(), q.order_index, q.question_text,
                        str(q.is_multiple_choice).lower(), ans.answer_text, ans.triage_level,
                    ]
                    ws.append(row)
                    triage_cell = ws.cell(row=ws.max_row, column=9)
                    triage_cell.fill = triage_fills.get(ans.triage_level, PatternFill())
            else:
                ws.append([
                    questionnaire.wound_type, questionnaire.title, questionnaire.description or "",
                    str(questionnaire.is_active).lower(), q.order_index, q.question_text,
                    str(q.is_multiple_choice).lower(), "", "green",
                ])

    col_widths = [12, 30, 30, 10, 14, 45, 18, 40, 12]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w

    # Add legend sheet only for single/bulk excel to assist users
    ws2 = wb.create_sheet("Hướng dẫn")
    ws2.append(["Cột", "Mô tả", "Giá trị hợp lệ"])
    ws2.append(["wound_type", "Loại vết thương", "burn, abrasion, bruise, fungal, laceration, rash, normal, cut, acne, psoriasis"])
    ws2.append(["title", "Tiêu đề bộ câu hỏi", "Chuỗi văn bản"])
    ws2.append(["description", "Mô tả (tùy chọn)", "Chuỗi văn bản hoặc để trống"])
    ws2.append(["is_active", "Trạng thái kích hoạt", "true / false"])
    ws2.append(["question_order", "Số thứ tự câu hỏi", "Số nguyên (1, 2, 3...)"])
    ws2.append(["question_text", "Nội dung câu hỏi", "Chuỗi văn bản"])
    ws2.append(["is_multiple_choice", "Chọn nhiều đáp án", "true / false"])
    ws2.append(["answer_text", "Nội dung đáp án", "Chuỗi văn bản"])
    ws2.append(["triage_level", "Mức độ nghiêm trọng", "green / yellow / red"])
    ws2.append([])
    ws2.append(["Lưu ý:", "Mẫu này export từ hệ thống SkinAid và có thể re-import."])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ─── Single CSV / Excel Export (Delegates to Bulk) ────────────────────────────

def export_to_csv(questionnaire: Questionnaire) -> bytes:
    """Export a single questionnaire to CSV by leveraging bulk export."""
    return export_bulk_to_csv([questionnaire])

def export_to_excel(questionnaire: Questionnaire) -> bytes:
    """Export a single questionnaire to Excel by leveraging bulk export."""
    return export_bulk_to_excel([questionnaire])
