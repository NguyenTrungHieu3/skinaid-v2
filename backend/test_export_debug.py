"""
Debug export - writes result to file to avoid encoding issues
"""
import io
import sys
import traceback

sys.stdout = open('test_output.txt', 'w', encoding='utf-8')
sys.stderr = sys.stdout

print(f"Python: {sys.version}")

# Test openpyxl
try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Test"
    headers = ["wound_type", "title", "question_text", "answer_text", "triage_level"]
    header_fill = PatternFill(fgColor="17805f", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.fill = header_fill
        cell.font = header_font
    ws.append(["burn", "Bo cau hoi bong", "Dien tich vet bong?", "Duoi 10cm2", "green"])
    buf = io.BytesIO()
    wb.save(buf)
    print(f"PASS openpyxl {openpyxl.__version__}: Excel size={len(buf.getvalue())} bytes")
except Exception as e:
    print(f"FAIL openpyxl: {e}")
    traceback.print_exc()

# Test CSV
try:
    import csv
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["wound_type", "title", "question_text", "answer_text", "triage_level"])
    writer.writerow(["burn", "Test", "Q1", "A1", "green"])
    result = output.getvalue().encode("utf-8-sig")
    print(f"PASS csv: {len(result)} bytes")
except Exception as e:
    print(f"FAIL csv: {e}")

# Test mock questionnaire export (mimic export_to_csv logic)
try:
    from dataclasses import dataclass, field
    from typing import List, Optional

    @dataclass
    class MockAnswer:
        answer_text: str
        triage_level: str
        order_index: int

    @dataclass
    class MockQuestion:
        question_text: str
        order_index: int
        is_multiple_choice: bool
        answers: List[MockAnswer]

    @dataclass
    class MockQuestionnaire:
        questionnaire_id: str = "test-uuid"
        wound_type: str = "burn"
        title: str = "Bo cau hoi bong"
        description: Optional[str] = "Mo ta"
        is_active: bool = True
        questions: List[MockQuestion] = field(default_factory=list)

    def mock_export_to_csv(questionnaire) -> bytes:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow([
            "wound_type", "title", "description", "is_active",
            "question_order", "question_text", "is_multiple_choice",
            "answer_text", "triage_level"
        ])
        questions = sorted(questionnaire.questions or [], key=lambda q: q.order_index)
        for q in questions:
            answers = sorted(q.answers or [], key=lambda a: a.order_index)
            if answers:
                for ans in answers:
                    writer.writerow([
                        questionnaire.wound_type,
                        questionnaire.title,
                        questionnaire.description or "",
                        str(questionnaire.is_active).lower(),
                        q.order_index,
                        q.question_text,
                        str(q.is_multiple_choice).lower(),
                        ans.answer_text,
                        ans.triage_level,
                    ])
        return output.getvalue().encode("utf-8-sig")

    q = MockQuestionnaire(questions=[
        MockQuestion("Dien tich vet bong?", 1, False, [
            MockAnswer("Duoi 10cm2", "green", 0),
            MockAnswer("Tu 10-50cm2", "yellow", 1),
            MockAnswer("Tren 50cm2", "red", 2),
        ]),
        MockQuestion("Vi tri vet bong?", 2, False, [
            MockAnswer("Tay hoac chan", "green", 0),
            MockAnswer("Mat hoac vung nhay cam", "red", 1),
        ]),
    ])
    csv_bytes = mock_export_to_csv(q)
    print(f"PASS mock CSV export: {len(csv_bytes)} bytes")
    print("CSV content sample:")
    print(csv_bytes.decode('utf-8-sig')[:300])

except Exception as e:
    print(f"FAIL mock export: {e}")
    traceback.print_exc()

# Test reportlab
try:
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Paragraph
    from reportlab.lib.styles import getSampleStyleSheet
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    styles = getSampleStyleSheet()
    story = [Paragraph("Test PDF", styles["Title"])]
    doc.build(story)
    import reportlab
    print(f"PASS reportlab {reportlab.Version}: PDF size={len(buf.getvalue())} bytes")
except Exception as e:
    print(f"FAIL reportlab: {e}")
    traceback.print_exc()

sys.stdout.flush()
